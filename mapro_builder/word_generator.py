from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
from typing import List, Dict

class MAPROWordGenerator:
    """Genera el documento Word del MAPRO"""
    
    def __init__(self, config: Dict, db_adapter):
        self.config = config
        self.db = db_adapter
        self.doc = Document()
        self.flujograma_gen = FlujogramaGenerator()
        self._configurar_estilos()
    
    def _configurar_estilos(self):
        """Configura estilos según el PDF base"""
        # Estilo Título 1
        style_h1 = self.doc.styles['Heading 1']
        style_h1.font.name = 'Arial'
        style_h1.font.size = Pt(14)
        style_h1.font.bold = True
        style_h1.font.color.rgb = RGBColor(0, 51, 102)
        
        # Estilo Título 2
        style_h2 = self.doc.styles['Heading 2']
        style_h2.font.name = 'Arial'
        style_h2.font.size = Pt(12)
        style_h2.font.bold = True
        
        # Estilo Normal
        style_normal = self.doc.styles['Normal']
        style_normal.font.name = 'Arial'
        style_normal.font.size = Pt(10)
    
    def generar_portada(self):
        """Genera portada del MAPRO"""
        # Encabezado institucional
        header = self.doc.sections[0].header
        header_para = header.paragraphs[0]
        header_para.text = f"MANUAL DE PROCEDIMIENTOS (MAPRO) – PAT. CLIN."
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Título principal
        self.doc.add_paragraph()
        titulo = self.doc.add_paragraph()
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = titulo.add_run(f"{self.config['establecimiento']['nombre']}")
        run.font.size = Pt(18)
        run.font.bold = True
        
        subtitulo = self.doc.add_paragraph()
        subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitulo.add_run(f"{self.config['establecimiento']['departamento']}")
        run.font.size = Pt(14)
        
        self.doc.add_paragraph()
        area = self.doc.add_paragraph()
        area.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = area.add_run(f"ÁREA DE {self.config['filtros']['upss_objetivo']}")
        run.font.size = Pt(16)
        run.font.bold = True
        
        self.doc.add_paragraph()
        anio = self.doc.add_paragraph()
        anio.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = anio.add_run(f"{self.config['establecimiento']['anio']}")
        run.font.size = Pt(14)
        
        # Tabla de aprobación
        self.doc.add_page_break()
        self._generar_tabla_aprobacion()
    
    def _generar_tabla_aprobacion(self):
        """Tabla Elaborado/Revisado/Aprobado"""
        tabla = self.doc.add_table(rows=4, cols=4)
        tabla.style = 'Table Grid'
        
        headers = ['ELABORADO POR', 'REVISADO POR', 'APROBADO POR', 'ÚLTIMA MODIFICACIÓN']
        for i, h in enumerate(headers):
            cell = tabla.rows[0].cells[i]
            cell.text = h
            cell.paragraphs[0].runs[0].font.bold = True
        
        tabla.rows[1].cells[0].text = self.config['establecimiento']['servicio']
        tabla.rows[1].cells[1].text = "OFIC. DE PLANEAMIENTO"
        tabla.rows[1].cells[2].text = "RES. DIRECT. N° ___"
        tabla.rows[1].cells[3].text = f"AÑO {self.config['establecimiento']['anio']}"
    
    def generar_capitulo_iv(self):
        """CAPÍTULO IV: Inventario de procedimientos"""
        self.doc.add_page_break()
        self.doc.add_heading('CAPÍTULO IV', level=1)
        self.doc.add_heading('INVENTARIO Y/O LISTADO DE PROCEDIMIENTOS', level=2)
        
        procedimientos = self.db.obtener_procedimientos()
        
        # Tabla de inventario
        tabla = self.doc.add_table(rows=len(procedimientos)+1, cols=5)
        tabla.style = 'Table Grid'
        
        headers = ['N°', 'CÓDIGO CPMS', 'DENOMINACIÓN', 'TIEMPO (min)', 'COSTO ESTÁNDAR']
        for i, h in enumerate(headers):
            cell = tabla.rows[0].cells[i]
            cell.text = h
            cell.paragraphs[0].runs[0].font.bold = True
        
        for idx, proc in enumerate(procedimientos, 1):
            tabla.rows[idx].cells[0].text = str(idx)
            tabla.rows[idx].cells[1].text = proc['codigo_cpms']
            tabla.rows[idx].cells[2].text = proc['nombre_procedimiento']
            tabla.rows[idx].cells[3].text = f"{proc['tiempo_total_minutos']:.1f}"
            tabla.rows[idx].cells[4].text = f"S/. {proc['costo_total_estandar']:.3f}"
    
    def generar_capitulo_v(self):
        """CAPÍTULO V: Fichas técnicas y flujogramas"""
        self.doc.add_page_break()
        self.doc.add_heading('CAPÍTULO V', level=1)
        self.doc.add_heading('FICHA DE DESCRIPCIÓN DE PROCEDIMIENTOS', level=2)
        
        procedimientos = self.db.obtener_procedimientos()
        
        for proc in procedimientos:
            self._generar_ficha_tecnica(proc)
            self.doc.add_page_break()
    
    def _generar_ficha_tecnica(self, procedimiento: Dict):
        """Genera una ficha técnica completa"""
        # Encabezado de ficha
        self.doc.add_heading(f"{procedimiento['nombre_procedimiento']}", level=2)
        
        # Tabla de metadatos
        tabla_meta = self.doc.add_table(rows=2, cols=4)
        tabla_meta.style = 'Table Grid'
        tabla_meta.rows[0].cells[0].text = "CÓDIGO CPMS:"
        tabla_meta.rows[0].cells[1].text = procedimiento['codigo_cpms']
        tabla_meta.rows[0].cells[2].text = "TIEMPO TOTAL:"
        tabla_meta.rows[0].cells[3].text = f"{procedimiento['tiempo_total_minutos']:.1f} min"
        tabla_meta.rows[1].cells[0].text = "COSTO ESTÁNDAR:"
        tabla_meta.rows[1].cells[1].text = f"S/. {procedimiento['costo_total_estandar']:.3f}"
        
        # Obtener fases
        fases = self.db.obtener_fases_procedimiento(procedimiento['id'])
        
        # OBJETIVO
        self.doc.add_heading('OBJETIVO', level=3)
        self.doc.add_paragraph(
            f"Establecer el procedimiento para la ejecución de "
            f"{procedimiento['nombre_procedimiento']}, asegurando resultados precisos y exactos."
        )
        
        # ALCANCE
        self.doc.add_heading('ALCANCE', level=3)
        self.doc.add_paragraph(
            f"Personal del área de {self.config['filtros']['upss_objetivo']}."
        )
        
        # MATERIALES Y EQUIPOS (agregados de todas las fases)
        self.doc.add_heading('MATERIALES Y EQUIPOS', level=3)
        
        # Recopilar equipos únicos
        equipos_unicos = set()
        for fase in fases:
            equipos = self.db.obtener_equipamiento_fase(fase['id'])
            for eq in equipos:
                equipos_unicos.add(eq['descripcion'])
        
        para = self.doc.add_paragraph()
        para.add_run('Equipos:').bold = True
        for eq in sorted(equipos_unicos):
            self.doc.add_paragraph(f"- {eq}", style='List Bullet')
        
        # Recopilar insumos por categoría
        insumos_por_categoria = {'EPP': set(), 'REACTIVO': set(), 'MATERIAL': set()}
        for fase in fases:
            insumos = self.db.obtener_insumos_fase(fase['id'])
            for cat, lista in insumos.items():
                for ins in lista:
                    insumos_por_categoria[cat].add(ins['descripcion'])
        
        para = self.doc.add_paragraph()
        para.add_run('Equipo de protección personal:').bold = True
        for epp in sorted(insumos_por_categoria['EPP']):
            self.doc.add_paragraph(f"- {epp}", style='List Bullet')
        
        # DESCRIPCIÓN DE PROCEDIMIENTOS (por fases)
        self.doc.add_heading('DESCRIPCIÓN DE PROCEDIMIENTOS', level=3)
        
        paso_num = 1
        for fase in fases:
            # Encabezado de fase
            para = self.doc.add_paragraph()
            run = para.add_run(f"FASE {fase['orden']}: {fase['nombre_fase']}")
            run.bold = True
            run.font.size = Pt(11)
            
            # Descripción de la actividad
            self.doc.add_paragraph(fase['descripcion_actividad'])
            
            # RRHH responsable
            rrhh = self.db.obtener_rrhh_fase(fase['id'])
            if rrhh:
                roles = [self.db.mapear_rol(r['grupo_ocupacional']) for r in rrhh]
                self.doc.add_paragraph(f"Responsable: {', '.join(roles)}")
            
            # Insumos específicos de esta fase
            insumos = self.db.obtener_insumos_fase(fase['id'])
            if insumos['REACTIVO']:
                para = self.doc.add_paragraph()
                para.add_run('Reactivos:').bold = True
                for ins in insumos['REACTIVO']:
                    self.doc.add_paragraph(
                        f"- {ins['descripcion']} ({ins['cantidad']} {ins['unidad_medida']})",
                        style='List Bullet'
                    )
            
            paso_num += 1
        
        # CRITERIOS DE RECHAZO (genérico, parametrizable)
        self.doc.add_heading('CRITERIOS DE RECHAZO', level=3)
        criterios = [
            "Muestras mal rotuladas o identificadas.",
            "Muestras con volumen insuficiente.",
            "Muestras hemolizadas o coaguladas (según aplique).",
            "Solicitud de laboratorio con datos incompletos."
        ]
        for crit in criterios:
            self.doc.add_paragraph(f"- {crit}", style='List Bullet')
        
        # FLUJOGRAMA
        self.doc.add_heading('FLUJOGRAMA', level=3)
        flujograma_path = self.flujograma_gen.generar_flujograma(procedimiento, fases)
        
        if flujograma_path.endswith('.png'):
            self.doc.add_picture(flujograma_path, width=Inches(6))
        else:
            self.doc.add_paragraph("[Código Mermaid generado - ver archivo adjunto]")
    
    def guardar(self, ruta_salida: str):
        """Guarda el documento Word"""
        os.makedirs(os.path.dirname(ruta_salida), exist_ok=True)
        self.doc.save(ruta_salida)
        return ruta_salida